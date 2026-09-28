"""Reorganize Proyect_US.pptx: insert CPS, sensitivity, dividers (+ countries
slide if present) as full-bleed picture slides, in a logical section order."""
import copy
from pptx import Presentation
from pptx.util import Emu
import os

SC = "/tmp/claude-0/-home-user-Alternative-Pathways----Quant/37b5ad96-b2b0-50d1-9e4b-6301ff28b789/scratchpad"
prs = Presentation(f"{SC}/deck.pptx")
W, H = prs.slide_width, prs.slide_height
blank = min(prs.slide_layouts, key=lambda l: len(l.placeholders))

def add_picture_slide(png):
    s = prs.slides.add_slide(blank)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    s.shapes.add_picture(png, 0, 0, width=W, height=H)
    return s

NEW = ["v_d1", "v_cps", "v_sens", "v_d2", "v_d3", "v_d4", "v_d5", "v_d6"]
have_countries = os.path.exists(f"{SC}/v_countries.png")
if have_countries:
    NEW.append("v_countries")
for tag in NEW:
    add_picture_slide(f"{SC}/{tag}.png")

# reorder: originals 0-26, new slides appended after (27..)
n = {tag: 27 + i for i, tag in enumerate(NEW)}
order = ([0, 1, n["v_d1"], n["v_cps"], 2, 3, n["v_sens"],
          n["v_d2"], 4, 5, 6, 7, 8, 9,
          n["v_d3"], 10, 11, 12, 13,
          n["v_d4"], 14, 15, 16, 17, 18,
          n["v_d5"], 19, 20,
          n["v_d6"], 21, 22, 23, 24, 25, 26]
         + ([n["v_countries"]] if have_countries else []))
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
assert len(ids) == len(order), (len(ids), len(order))
for el in ids:
    sldIdLst.remove(el)
for idx in order:
    sldIdLst.append(ids[idx])
prs.save(f"{SC}/deck_v2.pptx")
print(f"saved deck_v2.pptx with {len(order)} slides (countries: {have_countries})")
