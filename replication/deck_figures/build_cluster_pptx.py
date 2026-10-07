"""Native, editable PPTX mirror of the Data Cluster Beamer deck.
Writes report/data_cluster/Proyect_US_cluster.pptx. Run from the repo root.
Figures are taken as PNG from report/figures_deck; the equation image path
can be overridden with EQ_PNG."""
import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

FIG = "report/figures_deck"
OUT = "report/data_cluster/Proyect_US_cluster.pptx"
LOGO = "report/data_cluster/oecd_logo.png"
EQ_PNG = os.environ.get("EQ_PNG", "report/data_cluster/equation.png")

GARA = "Garamond"
INK = RGBColor(0x1A, 0x24, 0x30)
MUT = RGBColor(0x6B, 0x74, 0x80)
NAVY = RGBColor(0x1F, 0x38, 0x64)
HEAD = RGBColor(0x26, 0x45, 0x6E)
BAND = RGBColor(0x00, 0x54, 0x9F)
BURG = RGBColor(0x7B, 0x25, 0x30)
LGRAYB = RGBColor(0xC9, 0xCF, 0xD6)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def new_slide(band=True, title=None):
    s = prs.slides.add_slide(blank)
    if band:
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(-0.03),
                                Inches(13.34), Inches(0.65))
        sh.fill.solid(); sh.fill.fore_color.rgb = BAND
        sh.line.fill.background(); sh.shadow.inherit = False
    if title:
        box = s.shapes.add_textbox(Inches(0.65), Inches(0.82), Inches(12.0),
                                   Inches(0.6))
        tf = box.text_frame; tf.word_wrap = True; tf.auto_size = None
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = title
        r.font.name = GARA; r.font.size = Pt(25); r.font.bold = True
        r.font.color.rgb = HEAD
    return s

def add_text(s, x, y, w, h, parts, align=PP_ALIGN.LEFT, spacing=1.35):
    box = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, para in enumerate(parts):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = para[0] if isinstance(para[0], list) else [para]
        size = para[1] if not isinstance(para[0], list) else None
        if isinstance(para[0], list):
            p.space_after = Pt(para[1])
            for text, sz, color, bold, italic in para[0]:
                r = p.add_run(); r.text = text
                f = r.font
                f.name = GARA; f.size = Pt(sz); f.bold = bold
                f.italic = italic; f.color.rgb = color
            p.line_spacing = Pt(para[0][0][1] * spacing)
        else:
            text, size, color, bold, italic, after = para
            p.space_after = Pt(after)
            p.line_spacing = Pt(size * spacing)
            r = p.add_run(); r.text = text
            f = r.font
            f.name = GARA; f.size = Pt(size); f.bold = bold
            f.italic = italic; f.color.rgb = color
    return box

def rule(s, x, y, w, weight=1.0, color=INK):
    ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y),
                                Inches(x + w), Inches(y))
    ln.line.color.rgb = color; ln.line.width = Pt(weight)
    return ln

def fit_picture(s, path, box):
    bx, by, bw, bh = box
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(bw / iw, bh / ih)
    w, h = iw * scale, ih * scale
    s.shapes.add_picture(path, Inches(bx + (bw - w) / 2),
                         Inches(by + (bh - h) / 2), width=Inches(w))

def msg_policy(s, msg, pol, y=6.28):
    add_text(s, 0.7, y, 12.0, 0.45, [[
        [("Message.  ", 13, NAVY, True, False), (msg, 13, INK, False, False)], 0]])
    add_text(s, 0.7, y + 0.42, 12.0, 0.45, [[
        [("Policy.  ", 13, BURG, True, False), (pol, 13, INK, False, False)], 0]])

def chart_slide(title, png, msg, pol, box=(0.6, 1.5, 12.1, 4.65)):
    s = new_slide(title=title)
    fit_picture(s, f"{FIG}/{png}", box)
    msg_policy(s, msg, pol)
    return s

def cell(s, x, y, w, h, fill, border=None, label=None, lab_color=None):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill if fill else RGBColor(0xFF, 0xFF, 0xFF)
    if border is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = border; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    if label:
        tf = sh.text_frame; tf.word_wrap = False
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        r.font.name = GARA; r.font.size = Pt(13); r.font.bold = True
        r.font.color.rgb = lab_color or RGBColor(0xFF, 0xFF, 0xFF)
    return sh

# ================= 1. title =================
s = new_slide(band=True)
add_text(s, 1.2, 2.00, 10.93, 1.6, [
 ("Measuring Teacher Attrition with Labour Force Survey Data", 30, HEAD, True, False, 6),
 ("The Case of the United States", 20, HEAD, True, False, 0)], align=PP_ALIGN.CENTER)
add_text(s, 1.2, 4.05, 10.93, 0.5,
         [("José Elías Durán Roa", 16, INK, False, False, 0)],
         align=PP_ALIGN.CENTER)
add_text(s, 1.2, 4.60, 10.93, 0.4,
         [("Data Cluster · Directorate for Education and Skills", 12.5, MUT, False, False, 0)],
         align=PP_ALIGN.CENTER)
with Image.open(LOGO) as im:
    lw = 1.9; lh = lw * im.size[1] / im.size[0]
prs.slides[0].shapes.add_picture(LOGO, Inches((13.333 - lw) / 2), Inches(5.6),
                                 width=Inches(lw))

# ================= 2. why =================
s = new_slide(title="Why this project")
add_text(s, 0.9, 2.2, 11.5, 3.0, [
 ("•  Which of the factors we can observe correlate most strongly with "
  "the decision to leave teaching?", 18, INK, False, False, 22),
 ("•  Re-estimate the teacher leaving rate against the published figures "
  "of the official surveys.", 18, INK, False, False, 0)])

# ================= 3. contribution =================
s = new_slide(title="Contribution")
C = [("1.  A new series.  ",
      "The longest continuous series of teacher attrition to reach the "
      "present: one estimate for every year from 1997 to 2024, constructed "
      "from twenty-eight consecutive supplements of the Current Population "
      "Survey."),
     ("2.  A measurement discussion.  ",
      "A systematic treatment of the ways a longitudinal household survey "
      "identifies leaving: the linking of individuals across interviews, the "
      "alternative definitions of a leaver, and the assumptions each "
      "definition entails."),
     ("3.  Substantive evidence.  ",
      "A characterization of who leaves teaching and of the observed factors "
      "most strongly associated with leaving, together with the destinations "
      "of leavers and the heterogeneity of attrition across states.")]
y = 1.85
for lead, body in C:
    add_text(s, 0.9, y, 11.5, 1.4, [[
        [(lead, 16, NAVY, True, False), (body, 16, INK, False, False)], 0]],
        spacing=1.4)
    y += 1.55

# ================= 4. comparison table =================
s = new_slide(title="The CPS replicates the official picture, every year")
TX = [0.9, 4.7, 8.1, 10.5]
TW = [3.7, 3.1, 2.3, 2.0]
rule(s, 0.9, 1.70, 11.53, 1.4)
for x, w, htxt in zip(TX[1:], TW[1:], ["CPS", "NTPS (official)", "Difference"]):
    add_text(s, x, 1.80, w, 0.35, [(htxt, 14, INK, True, False, 0)],
             align=PP_ALIGN.RIGHT)
rule(s, 0.9, 2.22, 11.53, 0.75)
def trow(y, lab, a, b, d, bold=False, color=INK):
    add_text(s, TX[0], y, TW[0], 0.35, [(lab, 13.5, INK, bold, False, 0)])
    for x, w, v, c in zip(TX[1:], TW[1:], [a, b, d], [color, INK, INK]):
        if v:
            add_text(s, x, y, w, 0.35, [(v, 13.5, c, bold, False, 0)],
                     align=PP_ALIGN.RIGHT)
trow(2.34, "Measurement", "every year, 1997–2024", "occasional waves", "",
     bold=True, color=NAVY)
rule(s, 0.9, 2.78, 11.53, 0.75)
add_text(s, TX[0], 2.86, 8.0, 0.3,
         [("Example wave: 2020–21, the latest NTPS", 11.5, MUT, False, True, 0)])
trow(3.24, "Female (%)", "75.5", "76.8", "−1.3")
trow(3.66, "Master’s degree or higher (%)", "57.0", "61.0", "−4.0 ***")
trow(4.08, "Average age (years)", "44.0", "42.9", "+1.1 ***")
rule(s, 0.9, 4.52, 11.53, 0.75)
trow(4.62, "Teachers observed in the wave", "3,855", "39,633", "")
rule(s, 0.9, 5.04, 11.53, 1.4)
add_text(s, 0.9, 5.22, 11.53, 0.9, [
 ("Public school teachers. CPS: supplements of 2021–2022, weighted. NTPS 2020–21 "
  "(NCES 2022-113; Digest of Education Statistics, table 209.10; respondent count from "
  "NCES 2024-024). Stars test the CPS mean against the published NTPS estimate, "
  "design-adjusted standard errors: *** p<0.001; the female share does not differ "
  "significantly.", 11, MUT, False, True, 0)], align=PP_ALIGN.JUSTIFY)
add_text(s, 0.9, 6.35, 11.53, 0.8, [[
 [("Full sample:  ", 13.5, INK, False, False),
  ("5.0 million", 13.5, NAVY, True, False),
  (" person-year records, 1998–2025;  ", 13.5, INK, False, False),
  ("104,545", 13.5, NAVY, True, False),
  (" teacher observations;  ", 13.5, INK, False, False),
  ("48,842", 13.5, NAVY, True, False),
  (" teachers followed in the linked panel.", 13.5, INK, False, False)], 0]])

# ================= 5. rotation diagram (no title) =================
s = new_slide(band=True)
scw = 11.53 / 16
add_text(s, 0.9, 2.45, 3.5, 0.35, [("interviewed 4 months", 14, NAVY, True, False, 0)])
add_text(s, 4.9, 2.45, 3.6, 0.35, [("8 months out of the sample", 13, MUT, False, True, 0)],
         align=PP_ALIGN.CENTER)
add_text(s, 9.0, 2.45, 3.5, 0.35, [("interviewed 4 months", 14, BURG, True, False, 0)],
         align=PP_ALIGN.RIGHT)
for i in range(16):
    x0 = 0.9 + i * scw
    if i < 4:
        cell(s, x0, 2.90, scw, 0.62, NAVY, None, str(i + 1))
    elif i >= 12:
        cell(s, x0, 2.90, scw, 0.62, BURG, None, str(i + 5 - 12))
    else:
        cell(s, x0, 2.90, scw, 0.62, None, LGRAYB)
for xi, lab in [(0.9, "Mar"), (0.9 + 3 * scw, "Jun"),
                (0.9 + 12 * scw, "Mar"), (0.9 + 15 * scw, "Jun")]:
    add_text(s, xi, 3.58, scw, 0.3, [(lab, 11, MUT, False, False, 0)],
             align=PP_ALIGN.CENTER)
add_text(s, 0.9, 3.92, 4 * scw, 0.3, [("year 1", 12, MUT, False, True, 0)],
         align=PP_ALIGN.CENTER)
add_text(s, 0.9 + 12 * scw, 3.92, 4 * scw, 0.3, [("year 2", 12, MUT, False, True, 0)],
         align=PP_ALIGN.CENTER)
arrow = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(0.9 + scw / 2),
                               Inches(4.55), Inches(0.9 + 12.5 * scw), Inches(4.55))
arrow.line.color.rgb = INK; arrow.line.width = Pt(1.4)
add_text(s, 1.5, 4.72, 10.0, 0.4,
         [("the same interview, exactly twelve months later", 14, INK, False, False, 0)],
         align=PP_ALIGN.CENTER)

# ================= 6. ASEC =================
s = new_slide(title="The March supplement (ASEC)")
add_text(s, 0.9, 2.1, 11.5, 3.5, [
 ("•  The largest supplement of the CPS: the March sample is expanded to "
  "roughly 90,000 households.", 18, INK, False, False, 20),
 ("•  Every adult reports on the previous calendar year: longest job "
  "held, weeks worked, earnings, benefits.", 18, INK, False, False, 20),
 ("•  We use it to reconstruct the leaving rate.", 18, INK, False, False, 0)])

# ================= 7. codes =================
s = new_slide(band=False)
fit_picture(s, f"{FIG}/codes_slide.png", (0.3, 0.25, 12.73, 7.0))

# ================= 8. two ways =================
s = new_slide(title="Two ways to measure the leaving rate")
add_text(s, 0.9, 1.75, 5.6, 0.4, [("Follow the person, year to year", 16, NAVY, True, False, 0)])
add_text(s, 0.9, 2.35, 5.6, 2.2, [
 ("The rotation shows the same individual twelve months apart. The CPS has "
  "no person identifier, so records are linked on fields common to both "
  "years:", 14, INK, False, False, 10)], align=PP_ALIGN.JUSTIFY)
add_text(s, 0.9, 3.85, 5.6, 1.8, [
 ("household identifier", 14, INK, False, False, 6),
 ("dwelling identifier", 14, INK, False, False, 6),
 ("person line number within the household", 14, INK, False, False, 10),
 ("validated on sex, race, and age advancing 0–2 years", 13, MUT, False, False, 0)])
add_text(s, 7.0, 1.75, 5.5, 0.4, [("Ask about last year, each March", 16, NAVY, True, False, 0)])
add_text(s, 7.0, 2.35, 5.5, 2.6, [
 ("The supplement’s retrospective question. Whoever reports teaching as "
  "last year’s longest job, and no longer teaches at the March "
  "interview, left within the year.", 14, INK, False, False, 12),
 ("The measure used in the literature so far.", 13, MUT, False, False, 0)],
 align=PP_ALIGN.JUSTIFY)

# ================= 9. methodologies table =================
s = new_slide(title="Measuring the leaving rate: the alternatives, side by side")
MC = [(0.9, 1.9), (3.0, 3.1), (6.3, 3.6), (10.1, 2.4)]
rule(s, 0.9, 1.70, 11.6, 1.4)
for (x, w), htxt in zip(MC, ["Measure", "A teacher is…", "A leaver is…", "Rests on…"]):
    add_text(s, x, 1.80, w, 0.35, [(htxt, 13.5, INK, True, False, 0)])
rule(s, 0.9, 2.20, 11.6, 0.75)
MROWS = [
 ("Annual recall", "whoever taught as last year’s longest job",
  "no longer teaching at the March interview",
  "accurate recall of the previous year", INK, False),
 ("Occupation pair", "a teaching occupation code one year ago",
  "a different occupation code today",
  "a current code existing; fails for the non-employed", INK, False),
 ("Panel (benchmark)", "teaching in two or more year-1 interviews",
  "never teaching in year 2, given two or more interviews, one outside the summer",
  "correct linkage of persons across waves", NAVY, True),
 ("Month pairs, any sighting", "every teaching month, or a single sighting",
  "not teaching twelve months after each month",
  "one person counting several times; admits one-month spells", INK, False),
]
y = 2.36
for m, a, b, c, mc, mb in MROWS:
    add_text(s, MC[0][0], y, MC[0][1], 0.9, [(m, 12.5, mc, mb, False, 0)])
    add_text(s, MC[1][0], y, MC[1][1], 0.9, [(a, 12.5, INK, False, False, 0)])
    add_text(s, MC[2][0], y, MC[2][1], 0.9, [(b, 12.5, INK, False, False, 0)])
    add_text(s, MC[3][0], y, MC[3][1], 0.9, [(c, 12.5, INK, False, False, 0)])
    y += 1.02
rule(s, 0.9, y + 0.02, 11.6, 1.4)
add_text(s, 0.9, y + 0.18, 11.6, 0.4,
         [("Same records throughout; only the definition changes. Each row answers the same three questions.",
           12, MUT, False, True, 0)])

# ================= 10. all series =================
chart_slide("One dataset, one rate per definition", "v_allseries.png",
 "The definition of a leaver moves the measured rate by a factor of nearly four.",
 "Comparisons across studies and countries must hold the definition fixed.")

# ================= 11. framework =================
s = new_slide(title="Why teachers leave: what the CPS can and cannot see")
rule(s, 0.9, 1.80, 11.53, 1.4)
add_text(s, 1.1, 1.92, 7.0, 0.35, [("Reasons identified in the literature", 14, INK, True, False, 0)])
add_text(s, 9.3, 1.92, 3.0, 0.35, [("Observed in the CPS?", 14, INK, True, False, 0)])
rule(s, 0.9, 2.34, 11.53, 0.75)
FROWS = [("Working conditions", "no", MUT, False),
         ("Organizational support (mentoring, leadership)", "no", MUT, False),
         ("Autonomy and intrinsic motives", "no", MUT, False),
         ("Financial factors", "yes", NAVY, True),
         ("Personal life-cycle motives", "yes", NAVY, True)]
y = 2.48
for lab, v, c, b in FROWS:
    add_text(s, 1.1, y, 7.6, 0.35, [(lab, 14, INK, False, False, 0)])
    add_text(s, 9.3, y, 3.0, 0.35, [(v, 14, c, b, False, 0)])
    y += 0.48
rule(s, 0.9, y + 0.02, 11.53, 0.75)
add_text(s, 1.1, y + 0.12, 7.6, 0.35, [("Sociodemographic characteristics", 14, INK, False, False, 0)])
add_text(s, 9.3, y + 0.12, 3.0, 0.35, [("yes", 14, NAVY, True, False, 0)])
rule(s, 0.9, y + 0.58, 11.53, 1.4)
add_text(s, 0.9, y + 0.78, 11.53, 0.6,
         [("Five broad families of reasons from our review of the literature. The CPS covers the economic and life-cycle margins.",
           12, MUT, False, True, 0)])

# ================= 12. five messages =================
s = new_slide(title="Five messages")
FM = [
 "Economic and employment conditions correlate with leaving far more than demographics.",
 "Exits concentrate where attachment and annual pay are lowest: part-time and part-year teachers in the bottom of the earnings distribution.",
 "Those same exits lead out of employment, not to better-paid jobs: two thirds of bottom-quintile leavers exit the labor force. Low pay and non-employment destinations are one phenomenon, not two.",
 "Life-cycle events are the largest single trigger: childbirth raises a woman’s exit probability by about six points, more than in comparable professions.",
 "Pension coverage retains, and attrition differs widely across states: the actionable margins are contractual and institutional.",
]
y = 1.75
for i, m in enumerate(FM, 1):
    add_text(s, 0.9, y, 11.6, 1.0, [[
        [(f"{i}.  ", 15, NAVY, True, False), (m, 15, INK, False, False)], 0]],
        spacing=1.3)
    y += 1.12

# ================= 13+. chart run =================
RUN = [
 ("Leaving teaching, 1997–2024", "g2_evolution.png",
  "The rate fluctuates around 8.6 percent with no trend: a persistent outflow, not a mounting exodus.",
  "Retention faces a steady-state cost, not an emergency spike."),
 ("Leaving teaching, by school sector", "pubpriv_series.png",
  "Private schools lose teachers at roughly twice the public rate.",
  "Retention tracks contract structure, pay and pensions, not classroom conditions alone."),
 ("Leaving the occupation, by profession", "leave_professions_deck.png",
  "Teacher attrition is unexceptional next to comparable professions: above nurses and accountants, well below social workers.",
  "Benchmarks from other professions discipline how alarming any teacher figure should be."),
 ("Labor force exit of women, by profession", "lf_professions_deck.png",
  "Women teachers exit the labor force at about five percent a year, well above nurses, social workers and accountants.",
  "The family-compatibility margin is specific to teaching, not to professional work in general."),
 ("Teachers leaving per year, across countries", "intl_bars_deck.png",
  "The US sits in the upper range of the school systems that publish a rate.",
  "Comparable, survey-based measurement is what makes this ranking possible."),
 ("Destinations of teachers who leave, 2024", "g3_flow100_2024.png",
  "Of every hundred leavers, about a third take another job; most exit employment altogether.",
  "Competing employers are not the main counterfactual; retirement and care are."),
 ("The routes out of teaching, 2002–2024", "routes_bars3.png",
  "The destination mix has been essentially stable for two decades.",
  "The composition of the outflow is structural, not cyclical noise."),
 ("Who leaves the labor force before 55", "yellow_flow2024.png",
  "Nearly half of the pre-55 labor-force exits are mothers with a child at home.",
  "Childcare and parental-leave policy are retention policy."),
 ("Where the employed leavers go: every occupation", "occ_all_slide.png",
  "Employed leavers scatter widely; the largest block stays in education outside the classroom.",
  "Skills remain in the education orbit, so re-entry is a realistic margin."),
 ("Leaving teaching, by age and route", "g4_routes_age.png",
  "Leaving is U-shaped in age, and non-employment dominates both ends of the career.",
  "Early-career and pre-retirement exits call for different instruments."),
 ("EQUATION", None, None, None),
 ("Average marginal effects on the probability of leaving", "ame_seminar.png",
  "Job attachment and compensation dominate; demographic differences are small.",
  "The actionable levers are contractual: full-year contracts, pension coverage."),
 ("The age profile of leaving, by route", "u_routes_deck.png",
  "Young teachers switch to other jobs; older teachers leave the labor force.",
  "The two peaks of the U respond to different policies."),
 ("The effect of a new baby, by profession", "g8_newbaby.png",
  "A new baby raises a woman teacher’s exit probability by 5.9 points, the largest effect among comparable professions.",
  "Parental leave and childcare support target the largest single observable trigger."),
 ("Leaving teaching and the unemployment rate", "n5_cyclicality.png",
  "Exits into non-employment rise with unemployment; job-to-job exits do not.",
  "Downturns intensify the non-employment route; support should be countercyclical."),
 ("State-level cyclicality of leaving", "states_landscape.png",
  "States differ in how strongly teacher exits respond to the local cycle.",
  "National averages hide state-level exposure to downturns."),
 ("Attrition levels across states", "states_level.png",
  "Mean attrition of teachers under 55 ranges from about 4 to 13 percent across jurisdictions.",
  "State pay scales, pensions and institutions are first-order."),
 ("Real earnings growth by profession since 1997", "msgA3.png",
  "Teachers’ real median earnings fell 4 percent since 1997 while other graduates gained.",
  "The relative-pay gap compounds slowly; it is a stock, not a shock."),
 ("Leaving, by position in the teacher pay distribution", "pay3_deciles_deck.png",
  "Leaving concentrates sharply in the lowest earnings deciles, mostly part-year teachers.",
  "Stabilizing the bottom of the distribution buys the most retention."),
 ("Pay compression: teachers and comparable professions", "pay5_compression_deck.png",
  "Teacher pay is compressed: a 90/10 ratio of 2.6 against 5.1 for all other graduates.",
  "Staying offers little earnings upside; career ladders would widen the right tail."),
 ("The earnings change of leaving, by destination", "pay7_destinations_deck.png",
  "Leavers who move to other professional jobs gain about 12 log points; moves near education gain little.",
  "The outside option is real for those able to take it."),
 ("Earnings changes: stayers and leavers", "g5b_earnings_density.png",
  "For employed leavers the earnings change is a gamble: a quarter lose big, nearly half win big.",
  "Leaving buys risk, not a sure raise."),
 ("Pension coverage and leaving, across states", "pension_scatter_deck.png",
  "Where pension coverage is higher, fewer teachers under 55 leave: 0.21 points per point of coverage.",
  "Pension design is a measurable retention margin."),
]
for title, png, msg, pol in RUN:
    if title == "EQUATION":
        s = new_slide(title="The econometric specification")
        fit_picture(s, EQ_PNG, (1.2, 1.7, 10.9, 2.3))
        add_text(s, 0.9, 4.45, 11.6, 2.2, [
         ("•  Probit on college-graduate teachers, pooled 2015–2024, "
          "with calendar-year effects; n = 28,900.", 14.5, INK, False, False, 12),
         ("•  Average marginal effects reported, in percentage points.",
          14.5, INK, False, False, 12),
         ("•  Robust to logit and weighted LPM, state and year fixed "
          "effects, the full 1997–2024 window (n = 85,430), and the "
          "occupation-pair definition of leaving.", 14.5, INK, False, False, 0)])
    else:
        chart_slide(title, png, msg, pol)

# ================= conclusion =================
s = new_slide(title="Conclusion")
CO = [
 "A labour force survey can measure teacher attrition at the person level, yearly, over three decades, and it replicates the official picture of the workforce.",
 "The definition of a leaver is a first-order choice: the same records yield rates from 5 to 19 percent per year.",
 "What the data say: attachment, pay and pensions, and life-cycle events correlate with leaving; demographics much less.",
 "The design travels: the same construction is feasible in the labour force surveys of Italy, Germany, the UK, France, Spain, Brazil and Mexico.",
]
y = 1.9
for m in CO:
    add_text(s, 0.9, y, 11.6, 1.0,
             [("•  " + m, 16, INK, False, False, 0)], spacing=1.35)
    y += 1.15

prs.save(OUT)
print(f"saved {OUT} with {len(prs.slides._sldIdLst)} slides")
